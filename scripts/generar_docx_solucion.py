"""Genera docs/Solucion_Manaco_Data.docx con la explicación completa de la solución."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

BASE_DIR = Path(__file__).resolve().parent.parent
OUT_PATH = BASE_DIR / "docs" / "Solucion_Manaco_Data.docx"

ROJO = RGBColor(0xB3, 0x26, 0x1E)
GRIS = RGBColor(0x57, 0x53, 0x4E)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = color


def tabla(doc, encabezados, filas):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(encabezados):
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for fila in filas:
        row = t.add_row()
        for i, val in enumerate(fila):
            set_cell_text(row.cells[i], val)
    doc.add_paragraph()
    return t


def h1(doc, texto):
    p = doc.add_heading(texto, level=1)
    for run in p.runs:
        run.font.color.rgb = ROJO


def h2(doc, texto):
    p = doc.add_heading(texto, level=2)
    for run in p.runs:
        run.font.color.rgb = GRIS


def parrafo(doc, texto, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = bold
    run.italic = italic
    return p


def bullet(doc, texto):
    doc.add_paragraph(texto, style="List Bullet")


def numerado(doc, texto):
    doc.add_paragraph(texto, style="List Number")


def nota_caja(doc, titulo, texto):
    p = doc.add_paragraph()
    run = p.add_run(f"{titulo}: ")
    run.bold = True
    run.font.color.rgb = ROJO
    p.add_run(texto)


doc = Document()

estilo_normal = doc.styles["Normal"]
estilo_normal.font.name = "Calibri"
estilo_normal.font.size = Pt(11)

# --- Portada ------------------------------------------------------------
titulo = doc.add_heading("Manaco Data", level=0)
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in titulo.runs:
    run.font.color.rgb = ROJO

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run(
    "Sistema de control de inventario, mercado y pronóstico de demanda\npara Manaco (calzado, Bolivia)"
)
run.font.size = Pt(14)
run.font.color.rgb = GRIS

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Documento de explicación de la solución — caso Manaco").italic = True

doc.add_page_break()

# --- 1. Resumen ejecutivo ------------------------------------------------
h1(doc, "1. Resumen ejecutivo")
parrafo(
    doc,
    "Manaco Data es la solución construida para el caso Manaco: una aplicación web (Streamlit) con "
    "base de datos propia que resuelve los 4 problemas centrales identificados en el caso — falta de "
    "control diario de inventario en tienda, desconexión entre datos de mercado externo y ventas "
    "internas, falta de ajuste del pronóstico de demanda por presión del mercado informal, y falta de "
    "seguimiento formal de producción y calidad frente a la demanda proyectada.",
)
parrafo(
    doc,
    "El sistema cubre el ciclo completo: carga y limpieza de datos (ETL), registro diario de "
    "inventario con alertas automáticas, conteo asistido por visión artificial (IA), cruce de datos de "
    "comercio exterior con ventas internas, pronóstico de demanda ajustado por mercado informal, un "
    "prototipo de escenario para simular decisiones antes de tomarlas, y los 4 entregables escritos que "
    "exige la rúbrica del caso, ya redactados dentro del propio sistema.",
)

# --- 2. Arquitectura ------------------------------------------------------
h1(doc, "2. Arquitectura del sistema")
tabla(
    doc,
    ["Capa", "Tecnología", "Función"],
    [
        ("Interfaz", "Streamlit (multipágina)", "Páginas por rol, navegación con st.navigation + st.Page"),
        ("Lógica de negocio", "Python (src/services/)", "Reglas de inventario, ETL, pronóstico, visión por computadora"),
        ("Datos", "SQLite (manaco.db) / PostgreSQL opcional", "Catálogo, tiendas, ventas, cierres, comercio exterior"),
        ("Autenticación", "src/auth.py", "Login y control de acceso por rol (proteger_pagina)"),
        ("IA", "YOLO-World (yolov8s-world.pt)", "Conteo de pares de zapatos por foto en el cierre de tienda"),
    ],
)
parrafo(
    doc,
    "La base de datos usa SQLite por defecto (cero configuración) y puede apuntar a PostgreSQL "
    "configurando la variable de entorno DATABASE_URL, siguiendo la arquitectura propuesta en el caso.",
)

h2(doc, "2.1 Roles y permisos")
tabla(
    doc,
    ["Rol", "Puede ver"],
    [
        ("administrador", "Todo: catálogo, usuarios, carga de Excel, y las páginas de encargado y gerencia"),
        ("encargado", "Apertura, Ventas, Cierre y conteo (operación diaria de tienda)"),
        ("gerencia", "Inicio, Historial, Tablero de mercado, Pronóstico, Prototipo de escenario, Entregables escritos"),
    ],
)

# --- 3. Módulos funcionales ------------------------------------------------
h1(doc, "3. Módulos funcionales")

h2(doc, "3.1 Control diario de inventario")
parrafo(
    doc,
    "Cada tienda registra apertura (saldo inicial), ventas del día y cierre (conteo físico, manual o "
    "por foto con IA). El sistema calcula el saldo esperado (saldo inicial − vendido) y lo compara "
    "contra el conteo real, clasificando el resultado como Correcto, Faltante o Sobrante.",
)
nota_caja(
    doc,
    "Ejemplo de la demo",
    "\"Faltan 2 pares de MAN-DEP-001 en Tienda Prado (esperado 38, contado 36)\" — una diferencia que "
    "sin control diario solo se notaría en el cierre contable mensual.",
)

h2(doc, "3.2 Comercio exterior y tablero de mercado")
parrafo(
    doc,
    "El módulo de ETL (src/services/etl.py) limpia el Excel de comercio exterior de calzado "
    "(240 filas válidas; 20 comparten clave con otra fila y se marcan, sin descartarse), documentando "
    "cada regla de limpieza en una bitácora. El "
    "Tablero de mercado cruza esos datos con las ventas internas por categoría, usando el código del "
    "Sistema Armonizado (SA6) para enlazar el catálogo interno con las categorías del comercio exterior.",
)
bullet(doc, "Importación vs. exportación por categoría de calzado")
bullet(doc, "Top socios comerciales")
bullet(doc, "Valor por par (comparable entre categorías)")
parrafo(
    doc,
    "Hallazgo clave: la categoría Deportivo importa aproximadamente 3 veces lo que exporta, una señal "
    "de dependencia de oferta externa que antes era invisible para Manaco.",
)

h2(doc, "3.3 Pronóstico de demanda ajustado por mercado informal")
parrafo(
    doc,
    "El pronóstico de demanda por línea de producto (Deportivo, Caucho y cuero, Cuero) se ajusta por "
    "la presión del mercado informal/contrabando, usando una sensibilidad distinta por línea:",
)
tabla(
    doc,
    ["Línea de producto", "Sensibilidad al mercado informal"],
    [("Deportivo", "0.35 (la más expuesta)"), ("Caucho y cuero", "0.20"), ("Cuero", "0.10")],
)
parrafo(
    doc,
    "Un backtest compara el error del pronóstico ajustado contra el no ajustado (MAPE). Como el "
    "historial simulado ya incluye el efecto del mercado informal, el backtest demuestra que el método "
    "funciona, pero no mide la magnitud real del contrabando: eso se valida con ventas reales.",
)

h2(doc, "3.4 Prototipo de escenario")
parrafo(
    doc,
    "Permite simular, con controles deslizantes (tipo de cambio paralelo y presión informal), el "
    "impacto sobre 5 indicadores clave antes de tomar una decisión real:",
)
tabla(
    doc,
    ["Indicador", "Meta"],
    [
        ("Ventas proyectadas", "Referencia de planificación"),
        ("Cobertura de producción", "≥ 100% de la demanda proyectada"),
        ("Tasa de defectos de producción", "≤ 3%"),
        ("Brecha de precio vs. mercado informal", "≤ 20%"),
        ("Rechazo en ingesta de datos de sucursal", "≤ 3%"),
    ],
)

h2(doc, "3.5 Entregables escritos (dentro del propio sistema)")
parrafo(
    doc,
    "Los 4 textos que exige la rúbrica del caso están redactados y disponibles en vivo dentro de la "
    "aplicación, en la página \"Entregables escritos\" (solo administrador y gerencia), con 4 pestañas:",
)
numerado(doc, "Problemas prioritarios — tabla de 5 problemas con justificación y datos necesarios.")
numerado(doc, "Síntesis ejecutiva — 3 hallazgos, 2 limitaciones y 1 pregunta estratégica.")
numerado(doc, "Ficha de una página — problema, solución, 3 objetivos SMART/KPI, responsables, riesgos.")
numerado(doc, "Recomendación ejecutiva — comparación de 2 alternativas y 3 indicadores de seguimiento.")

# --- 4. Los 5 problemas priorizados --------------------------------------
h1(doc, "4. Problemas priorizados")
tabla(
    doc,
    ["N.º", "Problema y área", "Justificación"],
    [
        ("1", "Inventario sin control diario en tienda (producción, inventarios y logística)", "Sin conteo diario, una diferencia de stock solo se nota en el cierre contable mensual, cuando ya no se puede rastrear su causa."),
        ("2", "Presión de la oferta importada y del mercado informal (competitividad y mercado)", "Dato: Deportivo importa 2,97 veces lo que exporta. Hipótesis a validar: el calzado informal le quita ventas a Manaco."),
        ("3", "Rentabilidad por línea sin medir (rentabilidad y costos)", "El Cuero vale casi el doble por par en el mercado de referencia, pero Manaco no integra sus costos para saber qué línea le deja más utilidad."),
        ("4", "Cliente desconocido (clientes, segmentación y marketing)", "El sistema registra qué se vende, dónde y cuándo, pero no quién compra: no se puede segmentar ni medir el canal digital."),
        ("5", "Datos dispersos y sin responsables claros (gestión organizacional y calidad de datos)", "Mercado, ventas, inventario y calidad están separados, y los archivos de sucursal llegan con errores."),
    ],
)
parrafo(
    doc,
    "El Excel de comercio exterior responde comparaciones de categorías, socios y valor por par en un "
    "mercado de referencia; no responde costos, márgenes, clientes ni ventas de Manaco: eso exige datos internos.",
)

# --- 5. Síntesis ejecutiva -------------------------------------------------
h1(doc, "5. Síntesis ejecutiva del diagnóstico")
h2(doc, "3 hallazgos")
numerado(doc, "Deportivo es la categoría más dependiente de la importación: importa 2,97 veces lo que exporta en valor y 2,91 veces en pares. Como la proporción casi no cambia con la métrica, la conclusión es sólida.")
numerado(doc, "El calzado de cuero es un nicho de mayor valor por par: unos 41–45 USD por par frente a unos 25 USD de las otras dos categorías. Es pequeño en volumen, pero cada par vale casi el doble: es donde conviene diferenciarse por calidad y no por precio.")
numerado(doc, "Medir por valor o por cantidad cambia quiénes son los proveedores relevantes: por categoría el orden no cambia, pero República Checa entra al top 6 de proveedores por pares y no por valor, porque su calzado es el más barato de la base (unos 9 USD por par).")

h2(doc, "2 limitaciones")
numerado(doc, "La referencia es española y el valor es aduanero: sirve para comparar categorías y proporciones, pero no representa el mercado boliviano ni el precio al consumidor, que incluye aranceles, transporte y margen del comerciante.")
numerado(doc, "La cobertura por año es desigual (28 registros en 2020 y 82 en 2024): cualquier crecimiento entre años refleja la cantidad de registros, no el mercado, por lo que no se sacan tendencias temporales.")

h2(doc, "1 pregunta estratégica")
parrafo(
    doc,
    "Si Deportivo es la línea con más presión importada y Cuero la de mayor valor por par, ¿debe Manaco "
    "defender volumen en Deportivo o concentrar su diferenciación en Cuero? ¿Qué margen real por línea "
    "necesitaría ver para decidirlo?",
    italic=True,
)

h2(doc, "Respuestas a las 5 preguntas del caso")
tabla(
    doc,
    ["Pregunta del caso", "Respuesta con los datos"],
    [
        ('¿Qué datos internos debería recopilar Manaco?', 'Ventas por producto, tienda y día (ya las registra el sistema); costos reales por lote; defectos y devoluciones con su motivo; datos del cliente (segmento, ciudad y canal); fechas de entrega; inventario por almacén; y precios reales con descuentos.'),
        ('¿Qué hipótesis sobre competencia, abastecimiento o diferenciación podría investigar?', 'Competencia: Deportivo depende de la importación (importa 2,97 veces lo que exporta), así que es la línea más expuesta al calzado importado barato. Abastecimiento: hay proveedores muy baratos por par (República Checa, unos 9 USD) que sirven de referencia de costo. Diferenciación: el Cuero vale casi el doble por par, lo que sugiere un nicho de calidad.'),
        ('¿Cómo varía el valor por par y por qué no es el precio al consumidor?', 'Va de 8 a 119 USD por registro; la mediana es de unos 47 USD en Cuero y unos 27 USD en Deportivo y en Caucho y cuero. Es un valor aduanero: no incluye aranceles, impuestos, transporte ni el margen del comerciante, así que el precio en tienda es mayor.'),
        ('¿Qué países son proveedores y destinos relevantes?', 'Proveedores, por valor: Grecia, Países Bajos, Vietnam, Irlanda, Turquía y Rumanía. Destinos, por valor: Indonesia, Vietnam, Eslovenia, Irlanda, Grecia y Austria. Son relevantes dentro de los 30 socios del archivo, no del comercio mundial.'),
        ('¿Qué categorías tienen mayor valor? ¿Cambia al comparar con la cantidad?', 'Caucho y cuero lidera en ambos flujos (11.709,7 M USD importados y 8.514,2 M USD exportados), seguida de Deportivo; Cuero es muy pequeña. Por categoría el orden no cambia con la cantidad de pares; por socio sí: República Checa entra al top 6 de proveedores por pares, pero no por valor.'),
    ],
)

# --- 6. Ficha de una página -------------------------------------------------
h1(doc, "6. Ficha de propuesta estratégica")
tabla(
    doc,
    ["Campo", "Contenido"],
    [
        ("Problema", "Manaco no controla a diario su inventario ni mide la rentabilidad real de cada línea, y no cruza esos datos con el mercado."),
        ("Solución analítica", "Manaco Data: control diario con alertas, tablero de mercado cruzado por código SA6, y pronóstico con prototipo de escenario (hoy con datos [Simulado]). Apoya la decisión de qué línea priorizar y cuánto producir."),
        ("Datos que necesita", "Aperturas, ventas y cierres diarios (ya registrados); comercio exterior (ya limpio); costos, defectos y devoluciones por lote (por reemplazar con datos reales)."),
        ("Objetivo 1 — Inventario", "Cierres con diferencia por debajo del 10% al año 1, del 5% al año 2 y del 3% al año 3. KPI: % de cierres Correcto por tienda."),
        ("Objetivo 2 — Datos reales", "100% de tiendas con registro diario desde el mes 3; 24 meses de historia real al año 2; MAPE menor al 15% al año 3. KPI: % de tiendas con registro completo, meses de historia real y MAPE."),
        ("Objetivo 3 — Calidad", "Defectos de 3% o menos al año 1, 2,5% al año 2 y 2% al año 3; devoluciones por costura y pegado reducidas a la mitad. KPI: tasa de defectos y devoluciones por lote."),
        ("Responsables", "Gerencia General (comité mensual), Gerencia Comercial, Gerencia de Producción, encargados de tienda y Administración/TI."),
        ("Riesgo 1 y mitigación", "Registro diario incompleto: el cierre de inventario se suma al cierre de caja, con reporte semanal por tienda."),
        ("Riesgo 2 y mitigación", "Decidir con datos simulados: etiqueta [Simulado], sin comprometer presupuesto hasta tener 12 meses reales."),
        ("Viabilidad", "El sistema ya funciona en local; el costo principal es el tiempo del personal y la migración a PostgreSQL en el año 2."),
    ],
)

# --- 7. Recomendación ejecutiva --------------------------------------------
h1(doc, "7. Recomendación ejecutiva")
h2(doc, "Lectura de los indicadores")
bullet(doc, "Producción [Simulado]: lotes mensuales calibrados con la demanda base de cada línea (±15%).")
bullet(doc, "Defectos [Simulado]: entre 1% y 5% por lote; cerca de la mitad de los lotes supera la meta de 3%.")
bullet(doc, "Devoluciones [Simulado]: entre el 20% y el 50% de los pares defectuosos vuelven del cliente.")
bullet(doc, "Entregas: no hay datos; el sistema no registra fechas de despacho ni de entrega.")
bullet(doc, "Margen de catálogo [Supuesto]: cerca del 45% en las tres líneas, pero el Cuero deja unos 239 Bs por par frente a unos 155 Bs de las otras dos.")
parrafo(
    doc,
    "Falta para estimar rentabilidad: costos reales por lote, costos fijos, costo de cada devolución, "
    "gastos de almacenamiento y distribución, precio real con descuentos y fechas de entrega.",
)
tabla(
    doc,
    ["", "Alternativa A — Competir por precio", "Alternativa B — Diferenciar por calidad en Cuero"],
    [
        ("Acción", "Bajar el precio para cerrar la brecha informal.", "Mantener precio; control de calidad por lote y garantía en Cuero."),
        ("Beneficio", "Recuperar volumen.", "Menos devoluciones; marca más fuerte en la línea de mayor valor por par."),
        ("Costo", "Con margen de ~45%, bajar 10% el precio exige ~29% más pares para mantener la utilidad bruta.", "Inspección por lote, capacitación y stock de tallas; se cuantifica con costos reales."),
        ("Control", "Bajo (depende del tipo de cambio paralelo).", "Alto (procesos internos)."),
        ("Riesgo principal", "Guerra de precios contra un competidor sin costos formales.", "Que la garantía no se perciba si no se comunica."),
    ],
)
nota_caja(
    doc,
    "Recomendación: Alternativa B",
    "La brecha de precio depende del tipo de cambio paralelo, fuera del control de Manaco, y un 10% "
    "menos de precio exige casi un 30% más de volumen solo para quedar igual. La calidad es una palanca "
    "interna, y el Cuero es la línea donde el mercado paga más por par.",
)
h2(doc, "Segmento y acción comercial")
parrafo(
    doc,
    "Segmento [Supuesto]: adultos que trabajan en oficinas o servicios y compran calzado de cuero de vestir "
    "en las tiendas de La Paz, Cochabamba y Santa Cruz. El Cuero tiene el mayor valor por par, la menor "
    "presión importada (1,52 frente a 2,97 de Deportivo) y el mayor margen por par del catálogo.",
)
parrafo(
    doc,
    "Acción: lanzar \"Cuero garantizado\", con garantía de 6 meses en costura y suela, control de calidad "
    "de cada lote y stock asegurado de tallas centrales en las 3 tiendas.",
)
h2(doc, "3 indicadores de seguimiento")
numerado(doc, "Tasa de defectos y devoluciones de la línea Cuero (meta ≤3% de defectos).")
numerado(doc, "Ventas de la línea Cuero en pares y Bs frente a los 3 meses previos.")
numerado(doc, "% de cierres de inventario correctos por tienda.")

# --- 8. Limitaciones y próximos pasos --------------------------------------
h1(doc, "8. Limitaciones y próximos pasos")
parrafo(
    doc,
    "Los datos de comercio exterior son reales y ya están limpios en el sistema. Las ventas e inventario "
    "vienen de la demo, y los datos de producción, calidad, mercado local y precios de catálogo están "
    "explícitamente etiquetados [Simulado] o [Supuesto]. Antes de comprometer "
    "presupuesto real con los objetivos de este documento, deben reemplazarse por historia real de "
    "Manaco.",
)
bullet(doc, "Migrar de SQLite a PostgreSQL para soportar toda la red de tiendas.")
bullet(doc, "Recolectar al menos 12 meses de ventas reales para calibrar estacionalidad.")
bullet(doc, "Validar la sensibilidad al mercado informal con datos reales de precios de la competencia informal.")

doc.save(OUT_PATH)
print(f"Documento generado: {OUT_PATH}")
