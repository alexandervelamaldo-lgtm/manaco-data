"""Exportaciones a Excel con formato: planilla de cierres y tabla depurada de comercio exterior."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

RELLENO_ENCABEZADO = PatternFill("solid", fgColor="B3261E")
RELLENO_ESTADO = {
    "Faltante": PatternFill("solid", fgColor="FDECEA"),
    "Sobrante": PatternFill("solid", fgColor="FFF4E5"),
    "Correcto": PatternFill("solid", fgColor="EAF5EA"),
}
BORDE = Border(bottom=Side(style="thin", color="E7E5E4"))

COLUMNAS_EXPORTAR = [
    "Fecha", "Tienda", "Código", "Producto", "Categoría", "Saldo inicial", "Vendido",
    "Esperado", "Conteo IA", "Conteo final", "Diferencia", "Estado", "Confianza IA", "Observación",
]


def _escribir_tabla(hoja: Worksheet, df: pd.DataFrame, fila_inicio: int = 1) -> None:
    """Escribe un DataFrame con encabezados en negrita y anchos ajustados."""
    for j, col in enumerate(df.columns, start=1):
        celda = hoja.cell(row=fila_inicio, column=j, value=str(col))
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, fila in enumerate(df.itertuples(index=False), start=fila_inicio + 1):
        for j, valor in enumerate(fila, start=1):
            if not isinstance(valor, str) and pd.isna(valor):
                valor = None
            elif hasattr(valor, "item"):
                valor = valor.item()  # tipos numpy → tipos de Python
            celda = hoja.cell(row=i, column=j, value=valor)
            celda.border = BORDE
    for j, col in enumerate(df.columns, start=1):
        largos = [len(str(col))] + [len(str(v)) for v in df[col].head(500) if v is not None]
        hoja.column_dimensions[get_column_letter(j)].width = min(max(largos) + 3, 45)
    hoja.freeze_panes = hoja.cell(row=fila_inicio + 1, column=1)


def exportar_cierres_excel(cierres: pd.DataFrame) -> bytes:
    """Genera un Excel con la hoja ``Cierres`` (coloreada por estado) y una hoja ``Resumen``."""
    df = cierres.reindex(columns=COLUMNAS_EXPORTAR).copy()
    if "Confianza IA" in df:
        df["Confianza IA"] = df["Confianza IA"].astype(float).round(2)
    libro = Workbook()

    hoja = libro.active
    hoja.title = "Cierres"
    _escribir_tabla(hoja, df)
    col_estado = COLUMNAS_EXPORTAR.index("Estado") + 1
    col_fecha = COLUMNAS_EXPORTAR.index("Fecha") + 1
    for i in range(2, len(df) + 2):
        relleno = RELLENO_ESTADO.get(hoja.cell(row=i, column=col_estado).value)
        if relleno:
            for j in range(1, len(COLUMNAS_EXPORTAR) + 1):
                hoja.cell(row=i, column=j).fill = relleno
        hoja.cell(row=i, column=col_fecha).number_format = "DD/MM/YYYY"
    hoja.auto_filter.ref = hoja.dimensions

    resumen = libro.create_sheet("Resumen")
    resumen["A1"] = "Resumen de cierres de inventario — Manaco Data"
    resumen["A1"].font = Font(bold=True, size=14)
    if df.empty:
        resumen["A3"] = "No hay cierres con los filtros elegidos."
    else:
        resumen["A2"] = (f"Del {pd.to_datetime(df['Fecha']).min():%d/%m/%Y} "
                         f"al {pd.to_datetime(df['Fecha']).max():%d/%m/%Y} · {len(df)} cierres")
        por_estado = (df.groupby("Estado").agg(Cierres=("Estado", "size"), Diferencia_pares=("Diferencia", "sum"))
                      .reset_index().rename(columns={"Diferencia_pares": "Diferencia (pares)"}))
        _escribir_tabla(resumen, por_estado, fila_inicio=4)
        por_tienda = (df.groupby("Tienda").agg(
            Cierres=("Estado", "size"),
            Con_diferencia=("Diferencia", lambda s: int((s != 0).sum())),
            Faltantes=("Diferencia", lambda s: int(s[s < 0].sum())),
            Sobrantes=("Diferencia", lambda s: int(s[s > 0].sum())),
        ).reset_index().rename(columns={"Con_diferencia": "Con diferencia"}))
        inicio = 6 + len(por_estado) + 1
        _escribir_tabla(resumen, por_tienda, fila_inicio=inicio)
        resumen.freeze_panes = None
        resumen.column_dimensions["A"].width = 30

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def exportar_comercio_excel(datos: pd.DataFrame, diccionario: pd.DataFrame, bitacora: pd.DataFrame) -> bytes:
    """Excel de la tabla depurada de comercio exterior (Estación 2 del caso).

    ``datos`` usa los nombres de columna del Excel original y trae ``registro_valido`` y
    ``clave_repetida``. Hojas: ``Tabla depurada`` (registros válidos), ``Claves repetidas``
    (se conservan y se marcan para validarlas con la fuente), ``Diccionario`` y ``Bitácora``.
    """
    columnas = list(diccionario["Columna"])
    validos = datos[datos["registro_valido"]].copy()
    depurada = validos[columnas].copy()
    depurada["Clave repetida"] = validos["clave_repetida"].map({True: "Sí", False: "No"})
    repetidas = datos[datos["clave_repetida"]].sort_values(
        ["Flujo_comercial", "Año", "Código_SA6", "Socio_comercial"])[columnas]

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tabla depurada"
    _escribir_tabla(hoja, depurada)
    hoja.auto_filter.ref = hoja.dimensions

    hoja_rep = libro.create_sheet("Claves repetidas")
    hoja_rep["A1"] = ("Filas que comparten Flujo, Año, Código SA6 y Socio con otra fila, con valores distintos. "
                      "Se conservan y se marcan: deben validarse con la fuente original.")
    hoja_rep["A1"].font = Font(italic=True)
    _escribir_tabla(hoja_rep, repetidas, fila_inicio=3)

    hoja_dic = libro.create_sheet("Diccionario")
    _escribir_tabla(hoja_dic, diccionario)
    hoja_dic.column_dimensions["C"].width = 70
    for fila in hoja_dic.iter_rows(min_row=2, min_col=3, max_col=3):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")

    hoja_bit = libro.create_sheet("Bitácora")
    _escribir_tabla(hoja_bit, bitacora)
    hoja_bit.column_dimensions["C"].width = 80

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
