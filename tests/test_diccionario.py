"""Pruebas del diccionario de datos y de la exportación de la tabla depurada."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.config import CLAVE_REGISTRO, COLUMNAS_EXCEL, EXCEL_PATH
from src.services import diccionario as dic
from src.services import etl
from src.services.reportes import exportar_comercio_excel


@pytest.fixture(scope="module")
def datos() -> pd.DataFrame:
    df = etl.limpiar(etl.leer_excel(EXCEL_PATH)).datos
    df["clave_repetida"] = dic.marcar_claves_repetidas(df, CLAVE_REGISTRO)
    return df


def test_diccionario_cubre_las_10_columnas_sin_faltantes(datos):
    tabla = dic.construir_diccionario(datos)
    assert list(tabla["Columna"]) == list(COLUMNAS_EXCEL)
    assert tabla["Faltantes"].sum() == 0
    assert tabla["Descripción"].str.len().min() > 0


def test_diccionario_resume_rangos(datos):
    tabla = dic.construir_diccionario(datos).set_index("Columna")
    assert tabla.loc["Año", "Valores observados"] == "2020 a 2024"
    assert tabla.loc["País_reportante", "Valores observados"] == "España"
    assert tabla.loc["Flujo_comercial", "Valores observados"] == "Exportación · Importación"


def test_claves_repetidas_marca_ambas_filas(datos):
    assert int(datos["clave_repetida"].sum()) == 20


def test_exportar_tabla_depurada(datos):
    contenido = exportar_comercio_excel(datos, dic.construir_diccionario(datos),
                                        pd.DataFrame({"Regla": ["Resumen"], "Filas afectadas": [0],
                                                      "Detalle": ["240 → 240"]}))
    libro = load_workbook(BytesIO(contenido))
    assert libro.sheetnames == ["Tabla depurada", "Claves repetidas", "Diccionario", "Bitácora"]
    depurada = libro["Tabla depurada"]
    assert depurada.max_row - 1 == int(datos["registro_valido"].sum()) == 240
    encabezados = [c.value for c in depurada[1]]
    assert encabezados == list(COLUMNAS_EXCEL) + ["Clave repetida"]
    assert libro["Claves repetidas"].max_row - 3 == 20
    assert libro["Diccionario"].max_row - 1 == 10
